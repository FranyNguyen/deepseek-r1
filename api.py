from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from io import BytesIO
import requests
import openpyxl
from docx import Document

app = FastAPI()

VLLM_URL = "http://localhost:8000/v1/chat/completions"
stored_context = ""  # Store uploaded file content

def ask_vllm(question: str, context: str) -> str:
    """Send question to vLLM server and return response."""
    headers = {"Content-Type": "application/json"}
    data = {
        "model": "Qwen/Qwen2.5-1.5B-Instruct",
        "messages": [
            {
                "role": "system",
                "content": "Bạn là một chatbot thông minh. Trả lời hoàn toàn bằng ngôn ngữ tiếng Việt, ngắn gọn nhất"
            },
            {
                "role": "user",
                "content": f"Đây là nội dung bài viết:\n---\n{context}\n---\n\nDựa vào nội dung từ bài viết trên, hãy trả lời bằng ngôn ngữ tiếng Việt, ngắn gọn và đúng trọng tâm nhất.\n\n**Câu hỏi:** {question}"
            }
        ]
    }
    response = requests.post(VLLM_URL, headers=headers, json=data)
    if response.status_code != 200:
        return f"vLLM error: {response.text}"
    return response.json()["choices"][0]["message"]["content"]

@app.post("/upload/")
async def upload_file(file: UploadFile = File(...)):
    """Upload file, extract content, and store it."""
    global stored_context
    try:
        if file.filename.endswith(".txt"):
            stored_context = (await file.read()).decode("utf-8", errors="ignore")
        elif file.filename.endswith(".docx"):
            file_stream = BytesIO(await file.read())
            doc = Document(file_stream)
            stored_context = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        elif file.filename.endswith(".xlsx"):
            file_stream = BytesIO(await file.read())
            wb = openpyxl.load_workbook(file_stream, data_only=True)
            sheet = wb.active
            stored_context = "\n".join([" ".join([str(cell.value) if cell.value is not None else "" for cell in row]) for row in sheet.iter_rows()])
        else:
            raise HTTPException(status_code=400, detail="Only .txt, .docx, .xlsx supported")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"File processing error: {str(e)}")
    return {"message": "File uploaded successfully", "preview": stored_context[:500]}

@app.post("/ask/")
async def ask(question: str = Form(...)):
    """Answer question based on uploaded file content."""
    if not stored_context:
        raise HTTPException(status_code=400, detail="No file uploaded yet")
    answer = ask_vllm(question, stored_context)
    return {"answer": answer}